"""Тесты чтения корпоративных документов: DOCX/XLSX/PPTX (без ML-моделей)."""

import io
from datetime import datetime

import pytest

from docs import (
    SUPPORTED_EXT,
    read_document,
    read_docx,
    read_pdf,
    read_pptx,
    read_txt,
    read_xlsx,
)


# --- общий диспетчер ---

def test_supported_extensions():
    assert SUPPORTED_EXT == {".txt", ".pdf", ".docx", ".xlsx", ".pptx"}


def test_read_document_txt(tmp_path):
    p = tmp_path / "doc.txt"
    p.write_text("Привет, ГПН!", encoding="utf-8")
    assert "ГПН" in read_document(p)


def test_read_document_pdf(tmp_path):
    from pypdf import PdfWriter
    w = PdfWriter()
    w.add_blank_page(200, 200)
    p = tmp_path / "blank.pdf"
    with open(p, "wb") as f:
        w.write(f)
    assert read_document(p) == ""


def test_read_document_ext_override_stream(tmp_path):
    p = tmp_path / "файл"
    p.write_bytes("текст без расширения".encode("utf-8"))
    # BytesIO без имени — расширение передаётся явно
    with open(p, "rb") as f:
        text = read_document(io.BytesIO(f.read()), ext=".txt")
    assert text.strip() == "текст без расширения"


def test_read_document_unsupported_ext(tmp_path):
    p = tmp_path / "junk.xyz"
    p.write_bytes(b"x")
    with pytest.raises(ValueError):
        read_document(p)


def test_read_txt_binary_stream():
    assert read_txt(io.BytesIO("Русский текст".encode("utf-8"))) == "Русский текст"


# --- DOCX ---

def _make_docx(path):
    from docx import Document
    doc = Document()
    doc.add_heading("Стратегия", level=1)
    doc.add_heading("Цели 2025", level=2)
    doc.add_paragraph("Компания внедряет RPA и ML-решения для цифровизации.")
    doc.add_paragraph("Первый пункт инструкции.", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "Значение"
    table.cell(1, 0).text = "EBITDA"
    table.cell(1, 1).text = "120"
    doc.save(path)


def test_docx_heading_paragraph_and_table(tmp_path):
    p = tmp_path / "doc.docx"
    _make_docx(p)
    text = read_docx(p)

    # заголовки помечаются markdown и попадают в adaptive_chunk
    assert "# Стратегия" in text
    assert "# Цели 2025" in text
    # обычный параграф на месте
    assert "RPA и ML-решения" in text
    # таблица сериализуется построчно «ячейка | ячейка»
    assert "Показатель | Значение" in text
    assert "EBITDA | 120" in text


def test_docx_list_items_prefixed():
    p = io.BytesIO()
    _make_docx_to_stream(p)
    text = read_docx(p)
    assert "- Первый пункт инструкции." in text


def _make_docx_to_stream(stream):
    from docx import Document
    doc = Document()
    doc.add_paragraph("Первый пункт инструкции.", style="List Bullet")
    doc.save(stream)
    stream.seek(0)


def test_docx_from_stream(tmp_path):
    p = tmp_path / "stream.docx"
    _make_docx(p)
    with open(p, "rb") as f:
        text = read_docx(io.BytesIO(f.read()))
    assert "RPA и ML-решения" in text


def test_docx_corrupt_raises(tmp_path):
    p = tmp_path / "bad.docx"
    p.write_bytes(b"not a zip archive")
    with pytest.raises(Exception):
        read_docx(p)


# --- XLSX ---

def _make_xlsx(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Итоги"
    ws.append(["Показатель", "2024", "2025"])
    ws.append(["EBITDA", 120, 135])
    ws.append(["Выручка", 900.0, 950])
    wb.save(path)


def test_xlsx_header_aware_rows(tmp_path):
    p = tmp_path / "book.xlsx"
    _make_xlsx(p)
    text = read_xlsx(p)

    assert "[Лист: Итоги]" in text
    # строка данных обогащается заголовками: «Показатель: EBITDA | 2024: 120 | 2025: 135»
    assert "Показатель: EBITDA" in text
    assert "2024: 120" in text
    assert "2025: 135" in text
    # число 900.0 форматируется как целое
    assert "2024: 900" in text


def test_xlsx_single_column_list(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Термины"
    ws.append(["АЗС"])
    ws.append(["КРС"])
    wb.save(tmp_path / "list.xlsx")

    text = read_xlsx(tmp_path / "list.xlsx")
    lines = [line.strip() for line in text.strip().splitlines()]
    assert "[Лист: Термины]" in lines
    assert "АЗС" in lines
    assert "КРС" in lines


def test_xlsx_numeric_only_table_fallback_header(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    ws.append([120, 135])
    ws.append([900, 950])
    wb.save(tmp_path / "nohead.xlsx")

    text = read_xlsx(tmp_path / "nohead.xlsx")
    # без текстового заголовка первая строка из ≥2 ячеек считается шапкой
    assert "120: 900 | 135: 950" in text


def test_xlsx_dates_serialized(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Сроки"
    ws.append(["Дата", "Событие"])
    ws.append([datetime(2025, 1, 15, 10, 30), "Открытие центра"])
    wb.save(tmp_path / "dates.xlsx")

    text = read_xlsx(tmp_path / "dates.xlsx")
    assert "2025-01-15" in text
    assert "Открытие центра" in text


def test_xlsx_empty_workbook_returns_empty(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.save(tmp_path / "empty.xlsx")
    assert read_xlsx(tmp_path / "empty.xlsx") == ""


def test_xlsx_from_stream(tmp_path):
    p = tmp_path / "stream.xlsx"
    _make_xlsx(p)
    with open(p, "rb") as f:
        text = read_xlsx(io.BytesIO(f.read()))
    assert "2025: 135" in text


# --- PPTX ---

def _make_pptx(path):
    from pptx import Presentation
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tf = slide.shapes.add_textbox(0, 0, 9 * 914400, 2 * 914400).text_frame
    tf.paragraphs[0].text = "Ключевые результаты"
    tf.add_paragraph().text = "Добыча выросла на 5%"
    table_shape = slide.shapes.add_table(2, 2, 0, 3 * 914400, 9 * 914400, 2 * 914400)
    table = table_shape.table
    table.cell(0, 0).text = "Показатель"
    table.cell(0, 1).text = "2025"
    table.cell(1, 0).text = "EBITDA"
    table.cell(1, 1).text = "120"
    prs.save(path)


def test_pptx_text_and_tables(tmp_path):
    p = tmp_path / "deck.pptx"
    _make_pptx(p)
    text = read_pptx(p)

    assert "# Слайд 1" in text
    assert "Ключевые результаты" in text
    assert "Добыча выросла на 5%" in text
    assert "EBITDA | 120" in text


def test_pptx_from_stream(tmp_path):
    p = tmp_path / "stream.pptx"
    _make_pptx(p)
    with open(p, "rb") as f:
        text = read_pptx(io.BytesIO(f.read()))
    assert "# Слайд 1" in text


def test_pptx_empty_presentation_returns_empty(tmp_path):
    from pptx import Presentation
    p = tmp_path / "empty.pptx"
    Presentation().save(p)
    assert read_pptx(p) == ""